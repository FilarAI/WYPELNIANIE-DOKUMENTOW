"""Decyzje, teczki i PDF są źródłem kandydatów, nie ślepą kopią liczb."""
import io
import re

import fitz
from docx import Document

from app.main import app
from app.odczyt_plikow import odczytaj_tekst
from app.scalanie_zrodel import odczytaj_zestaw
from app.generator import generuj_dokumenty


def decyzja(imiona="Jan i Anna Kowalscy", dzialka="1358", znak="AB.123.2025"):
    d = Document()
    for linia in (
        "STAROSTA NOWOTARSKI", f"Nasz znak: {znak}", "DECYZJA NR 45/2025",
        "Nowy Targ, dnia 02.04.2025 r.",
        "ZATWIERDZAM PROJEKT BUDOWLANY I UDZIELAM POZWOLENIA NA BUDOWĘ",
        f"dla: Państwa {imiona}, zam. 32-400 Myślenice 4,",
        "OBEJMUJĄCE:", "budowę budynku mieszkalnego jednorodzinnego",
        f"na działce ewid. nr {dzialka}, położonej w miejscowości Myślenice, obręb 0002 Myślenice, jednostka ewidencyjna 120903_4 Myślenice,",
        "dane techniczne budynku:", "- kubatura", "- pow. zabudowy", "520,10 m3 110,20 m2",
        "- pow. użytkowa", "- ilość izb mieszkalnych", "- ilość lokali mieszkalnych",
        "100,30 m2", "5", "1", "I", "UZASADNIENIE",
        "W dniu 04.04.2025 r. Inwestorzy, tj. Państwo Jan i Anna Kowalscy, wystąpili z wnioskiem.",
        "decyzja stała się ostateczna z dniem 20.04.2025 r.",
    ):
        d.add_paragraph(linia)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_decyzja_wydobywa_parametry_ale_nie_miesza_dat():
    dane = odczytaj_zestaw([("decyzja", "decyzja.docx", decyzja())])
    assert dane["zakonczenie"]["pozwolenie_data"] == "02.04.2025"
    assert dane["zakonczenie"]["pozwolenie_znak"] == "AB.123.2025"
    assert dane["zakonczenie"]["kubatura_mieszkalna"] == "520,10"
    assert dane["zakonczenie"]["pow_zabudowy_mieszkalna"] == "110,20"
    assert dane["zakonczenie"]["pow_uzytkowa_mieszkalna"] == "100,30"
    assert dane["zakonczenie"]["liczba_izb"] == "5"
    assert dane["zakonczenie"]["liczba_mieszkan"] == "1"
    assert dane["adres_inwestycji"]["identyfikator_dzialki"] == "120903_4.0002.1358"
    assert not dane["blokada_zestawu"]


def test_teczka_nie_bierze_pojemnosci_gazu_jako_sciekow():
    d = Document(io.BytesIO(decyzja()))
    for linia in ("PROTOKÓŁ ODBIORCZY", "INSTALACJI C.O. WRAZ Z KOTŁEM GAZOWYM",
                  "Podziemny zbiornik gazowy o pojemności 10 m3",
                  "PROTOKÓŁ ODBIORCZY", "PRZYŁĄCZA KANALIZACJI SANITARNEJ DO PRZYDOMOWEJ",
                  "OCZYSZCZALNI ŚCIEKÓW", "Oczyszczalnia nadaje się do użytkowania."):
        d.add_paragraph(linia)
    b = io.BytesIO()
    d.save(b)
    dane = odczytaj_zestaw([("teczka", "teczka.docx", b.getvalue())])
    assert "ogrzewanie" not in dane["zakonczenie"]
    assert "kanalizacja" not in dane["zakonczenie"]
    assert {p["wartosc"] for p in dane["potwierdzenia"]} == {"gazowe", "oczyszczalnia"}
    assert "zbiornik_pojemnosc" not in dane["zakonczenie"]
    assert "gaz_sieciowy" not in dane["zakonczenie"]


def test_rozne_inwestycje_blokuja_generowanie():
    a = decyzja()
    b = decyzja(imiona="Piotr i Maria Nowak", dzialka="999", znak="BA.999.2025")
    dane = odczytaj_zestaw([("decyzja", "a.docx", a), ("decyzja", "b.docx", b)])
    assert dane["blokada_zestawu"]
    app.config["TESTING"] = True
    with app.test_client() as klient:
        odpowiedz = klient.post("/wybierz-dokumenty", data={"typ_0": "decyzja", "obiegowka": (io.BytesIO(a), "a.docx"),
            "typ_1": "decyzja", "plik_1": (io.BytesIO(b), "b.docx")}, content_type="multipart/form-data")
        token = re.search(r'name="token" value="([^"]+)"', odpowiedz.get_data(as_text=True)).group(1)
        podglad = klient.post("/podglad", data={"token": token, "dokumenty": "zawiadomienie"})
        assert podglad.status_code == 200
        assert "Nie można połączyć tych plików" in podglad.get_data(as_text=True)
        assert klient.post("/generuj", data={"token_wgrania": token, "inw-0-imie_nazwisko": "Jan Kowalski"}).status_code == 400


def test_pdf_z_warstwa_tekstowa():
    pdf = fitz.open()
    strona = pdf.new_page()
    strona.insert_text((50, 70), "DECYZJA NR 45/2025", fontsize=14)
    strona.insert_text((50, 100), "Nasz znak: AB.123.2025", fontsize=14)
    linie, ocr = odczytaj_tekst(pdf.tobytes(), "decyzja.pdf")
    assert not ocr
    assert any("DECYZJA NR 45/2025" in tekst for _, tekst in linie)


def test_wybranie_zlego_typu_plikow():
    try:
        odczytaj_zestaw([("obiegowka", "decyzja.docx", decyzja())])
    except ValueError as e:
        assert "Nie znaleziono" in str(e)
    else:
        assert False, "Decyzja nie może być odczytana jako obiegówka"


def test_wiele_dzialek_wymaga_wyboru_a_pb5_ma_zalacznik():
    d = Document(io.BytesIO(decyzja()))
    for akapit in d.paragraphs:
        if akapit.text.startswith("na działce ewid."):
            akapit.text = ("zjazd z drogi gminnej - dz. ewid. ozn. nr 3507) na działkach ewid. nr "
                            "1358, 1359, 3507, położonych w miejscowości Myślenice, "
                            "obręb 0002 Myślenice, jednostka ewidencyjna 120903_4 Myślenice,")
    b = io.BytesIO()
    d.save(b)
    dane = odczytaj_zestaw([("decyzja", "decyzja.docx", b.getvalue())])
    assert dane["adres_inwestycji"]["nr_dzialki"] == ""
    assert len(dane["dzialki_do_wyboru"]) == 3
    assert dane["dzialki_do_wyboru"][-1]["droga"]
    wybrane = dane["dzialki_do_wyboru"][:2]
    dane["adres_inwestycji"]["nr_dzialki"] = ", ".join(x["numer"] for x in wybrane)
    dane["adres_inwestycji"]["identyfikator_dzialki"] = ", ".join(x["identyfikator"] for x in wybrane)
    dane["dokumenty"] = ("pb5",)
    plik = next(iter(generuj_dokumenty(dane).values()))
    wynik = Document(io.BytesIO(plik))
    akapity = "\n".join(x.text for x in wynik.paragraphs)
    tabele = "\n".join(c.text for t in wynik.tables for row in t.rows for c in row.cells)
    assert "Liczba stron zawierających dane o kolejnych nieruchomościach (załączanych do oświadczenia): 1" in akapity
    assert "120903_4.0002.1358" in akapity
    assert "120903_4.0002.1359" in tabele
    assert "120903_4.0002.3507" not in akapity + tabele


def test_teczka_odczytuje_decyzje_za_strona_tytulowa():
    d = Document(io.BytesIO(decyzja()))
    for n in range(210):
        d.paragraphs[0].insert_paragraph_before(f"Spis załączników {n}")
    b = io.BytesIO()
    d.save(b)
    dane = odczytaj_zestaw([("teczka", "teczka.docx", b.getvalue())])
    assert dane["zakonczenie"]["pozwolenie_znak"] == "AB.123.2025"


def test_rozne_miejscowosci_blokuja_nawet_przy_zgodnym_znaku():
    a = decyzja()
    d = Document(io.BytesIO(decyzja()))
    for akapit in d.paragraphs:
        if "położonej w miejscowości Myślenice" in akapit.text:
            akapit.text = akapit.text.replace("położonej w miejscowości Myślenice",
                                                "położonej w miejscowości Dobczyce")
    b = io.BytesIO()
    d.save(b)
    dane = odczytaj_zestaw([("decyzja", "a.docx", a), ("decyzja", "b.docx", b.getvalue())])
    assert dane["blokada_zestawu"]
