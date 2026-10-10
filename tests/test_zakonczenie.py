"""Wybór i dane do zakończenia budowy bez zgadywania oświadczeń."""
import io
import re
import zipfile
from pathlib import Path

from docx import Document
from werkzeug.datastructures import MultiDict

from app.main import _dane_z_formularza, app


FIXTURES = Path(__file__).parent / "fixtures"


def test_podglad_pokazuje_pola_i_nie_zaznacza_oswiadczen():
    app.config["TESTING"] = True
    with app.test_client() as klient, open(FIXTURES / "obiegowka_przyklad.docx", "rb") as plik:
        wybor = klient.post("/wybierz-dokumenty", data={"obiegowka": (plik, "obiegowka.docx")},
                            content_type="multipart/form-data")
        token = re.search(r'name="token" value="([^"]+)"', wybor.get_data(as_text=True)).group(1)
        odpowiedz = klient.post("/podglad", data=MultiDict([
            ("token", token), ("dokumenty", "oswiadczenie"), ("dokumenty", "zawiadomienie")]))
    html = odpowiedz.get_data(as_text=True)
    assert odpowiedz.status_code == 200
    assert "Zakończenie budowy" in html
    assert "Bartłomiej Filar" in html
    assert 'name="zak-kierownik_imie_nazwisko"' not in html
    assert 'name="zak-pozwolenie_znak"' in html
    assert 'name="zak-zalaczniki"' in html
    assert 'name="data_pelnomocnictwa"' not in html
    assert '<select name="zak-zgodnosc">' in html
    assert '<option value="">— wybierz —</option>' in html
    assert 'name="zak-rodzaje_obiektu" value="nowy" checked' not in html


def test_dane_z_formularza_zachowuja_wybrane_wartosci():
    formularz = MultiDict([
        ("inw-0-imie_nazwisko", "Jan Kowalski"),
        ("zak-kierownik_imie_nazwisko", "Anna Nowak"),
        ("zak-ep", "72,5"),
        ("zak-zgodnosc", "odstepstwa"),
        ("zak-odstepstwa", "Zmiana położenia okna"),
        ("zak-rodzaje_obiektu", "nadbudowa"),
        ("zak-rodzaje_obiektu", "mieszkalny"),
        ("zak-zalaczniki", "dziennik_budowy"),
    ])
    dane = _dane_z_formularza(formularz)
    zak = dane["zakonczenie"]
    assert zak["kierownik_imie_nazwisko"] == "Bartłomiej Filar"
    assert zak["ep"] == "72,5"
    assert zak["zgodnosc"] == "odstepstwa"
    assert zak["odstepstwa"] == "Zmiana położenia okna"
    assert zak["rodzaje_obiektu"] == ["nadbudowa", "mieszkalny"]
    assert zak["zalaczniki"] == ["dziennik_budowy"]
    assert zak["obiekt_nadaje_sie"] == ""


def test_niepoprawny_wybor_jest_odrzucany():
    app.config["TESTING"] = True
    with app.test_client() as klient:
        odpowiedz = klient.post("/generuj", data={
            "inw-0-imie_nazwisko": "Jan Kowalski",
            "zak-zgodnosc": "dowolna wartość",
        })
    assert odpowiedz.status_code == 400


def test_wybor_jednego_dokumentu_i_wypelnienie_wzoru():
    app.config["TESTING"] = True
    with app.test_client() as klient:
        odpowiedz = klient.post("/generuj", data={
            "dokumenty": "oswiadczenie",
            "inw-0-imie_nazwisko": "Jan Kowalski",
            "nazwa_inwestycji": "BUDOWA DOMU",
            "adres-miejscowosc": "Myślenice", "adres-nr_dzialki": "1358",
            "zak-data_sporzadzenia": "10.10.2026", "zak-miejscowosc_sporzadzenia": "Rabka-Zdrój",
            "zak-pozwolenie_znak": "AB.123.2025", "zak-zgodnosc": "zgodnie",
            "zak-u_sciany": "0,20",
            "zak-kierownik_imie_nazwisko": "Podszywanie się",
        })
    assert odpowiedz.status_code == 200
    with zipfile.ZipFile(io.BytesIO(odpowiedz.data)) as archiwum:
        nazwy = archiwum.namelist()
        assert nazwy == ["Oswiadczenie_kierownika_Kowalski.docx"]
        doc = Document(io.BytesIO(archiwum.read(nazwy[0])))
    tekst = "\n".join(p.text for p in doc.paragraphs)
    assert "Bartłomiej Filar" in tekst
    assert "Podszywanie się" not in tekst
    assert "BUDOWA DOMU" in tekst and "AB.123.2025" in tekst
    assert "0,20" in tekst and "Myślenice" in tekst
    assert "[WYBRANO]" in tekst


def test_trzy_wybrane_dokumenty_bez_pelnomocnictwa():
    app.config["TESTING"] = True
    with app.test_client() as klient:
        odpowiedz = klient.post("/generuj", data=MultiDict([
            ("dokumenty", "pb5"), ("dokumenty", "oswiadczenie"),
            ("dokumenty", "zawiadomienie"), ("inw-0-imie_nazwisko", "Jan Kowalski"),
            ("zak-kubatura_mieszkalna", "500"), ("zak-rodzaje_obiektu", "nowy"),
            ("zak-zalaczniki", "dziennik_budowy")]))
    assert odpowiedz.status_code == 200
    with zipfile.ZipFile(io.BytesIO(odpowiedz.data)) as archiwum:
        nazwy = archiwum.namelist()
        tekst = "\n".join(p.text for p in Document(io.BytesIO(
            archiwum.read(next(n for n in nazwy if n.startswith("Zawiadomienie"))))).paragraphs)
    assert len(nazwy) == 3
    assert not any(n.startswith("Pelnomocnictwo") for n in nazwy)
    assert "500" in tekst and "[x]" in tekst


def test_niepoprawny_lub_pusty_wybor_dokumentow():
    app.config["TESTING"] = True
    with app.test_client() as klient:
        assert klient.post("/generuj", data={"dokumenty": "nieznany",
                                              "inw-0-imie_nazwisko": "Jan Kowalski"}).status_code == 400
        assert klient.post("/generuj", data={"wybor_potwierdzony": "tak",
                                              "inw-0-imie_nazwisko": "Jan Kowalski"}).status_code == 400
