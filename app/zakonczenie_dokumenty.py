"""Wypełnianie dwóch wzorów zakończenia budowy dostarczonych przez użytkownika.

Zmiany dotyczą wyłącznie miejsc do wpisania danych i oznaczenia wariantów.
Podpisów ani niepotwierdzonych oświadczeń program nie składa za człowieka.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document

SZABLONY = Path(__file__).resolve().parent.parent / "szablony"
KROPKI = re.compile(r"[.…]{2,}")


def _wymien_fragment(akapit, poczatek: int, koniec: int, wartosc: str) -> None:
    """Podmień fragment przez runy bez usuwania pozostałego formatowania."""
    przebieg = 0
    dotkniete = []
    for run in akapit.runs:
        koniec_runu = przebieg + len(run.text)
        if przebieg < koniec and koniec_runu > poczatek:
            dotkniete.append((run, max(0, poczatek - przebieg), min(len(run.text), koniec - przebieg)))
        przebieg = koniec_runu
    if not dotkniete:
        return
    pierwszy, od, do = dotkniete[0]
    if len(dotkniete) == 1:
        pierwszy.text = pierwszy.text[:od] + wartosc + pierwszy.text[do:]
        return
    pierwszy.text = pierwszy.text[:od] + wartosc
    for run, _, _ in dotkniete[1:-1]:
        run.text = ""
    ostatni, _, do = dotkniete[-1]
    ostatni.text = ostatni.text[do:]


def _kropki(akapit, *wartosci: str) -> None:
    oryginal = akapit.text
    miejsca = list(KROPKI.finditer(oryginal))
    if len(wartosci) > len(miejsca):
        raise ValueError("Zmieniono układ wzoru dokumentu: brakuje pól do wypełnienia.")
    for i in range(len(wartosci) - 1, -1, -1):
        wartosc = str(wartosci[i] or "").strip()
        if wartosc:
            przed = oryginal[miejsca[i].start() - 1:miejsca[i].start()]
            po = oryginal[miejsca[i].end():miejsca[i].end() + 1]
            if przed and not przed.isspace() and przed not in "(—":
                wartosc = " " + wartosc
            if po and not po.isspace() and po not in ",.;:)]":
                wartosc += " "
            _wymien_fragment(akapit, miejsca[i].start(), miejsca[i].end(), wartosc)


def _zamien(akapit, stary: str, nowy: str) -> None:
    pozycja = akapit.text.find(stary)
    if pozycja < 0:
        raise ValueError("Zmieniono tekst wzoru dokumentu: " + stary)
    _wymien_fragment(akapit, pozycja, pozycja + len(stary), nowy)


def _dopisz(akapit, tekst: str) -> None:
    akapit.add_run(tekst)


def _oznacz(akapit, wybrane: list[str] | str, mapa: dict[str, str]) -> None:
    """Podkreślenie odpowiedzi, jak nakazuje oryginalny druk."""
    if isinstance(wybrane, str):
        wybrane = [wybrane] if wybrane else []
    fragmenty = [(akapit.text.rfind(mapa[k]) if k == "sprzedaz_wynajem" else akapit.text.find(mapa[k]), mapa[k])
                 for k in wybrane if k in mapa]
    # Odtworzenie tylko akapitów z listami wariantów, z zachowaniem stylu akapitu.
    tekst = akapit.text
    zakresy = sorted((start, start + len(fragment)) for start, fragment in fragmenty if start >= 0)
    if not zakresy:
        return
    for run in akapit.runs:
        run.text = ""
    pozycja = 0
    for poczatek, koniec in zakresy:
        if poczatek < pozycja:
            continue
        akapit.add_run(tekst[pozycja:poczatek])
        run = akapit.add_run(tekst[poczatek:koniec])
        run.underline = True
        run.bold = True
        pozycja = koniec
    akapit.add_run(tekst[pozycja:])


def _adres(i: dict) -> str:
    czesc = " ".join(x for x in (i.get("ulica", ""), i.get("nr_domu", "")) if x)
    if i.get("nr_lokalu"):
        czesc += "/" + i["nr_lokalu"]
    druga = " ".join(x for x in (i.get("kod_pocztowy", ""), i.get("miejscowosc", "")) if x)
    return ", ".join(x for x in (czesc, druga) if x)


def _inwestorzy(dane: dict) -> str:
    return ", ".join(i["imie_nazwisko"] for i in dane["inwestorzy"] if i.get("imie_nazwisko"))


def _zapisz(dokument) -> bytes:
    bufor = io.BytesIO()
    dokument.save(bufor)
    return bufor.getvalue()


def oswiadczenie(dane: dict) -> bytes:
    z = dane.get("zakonczenie", {})
    adres = dane.get("adres_inwestycji", {})
    p = Document(SZABLONY / "oswiadczenie_kierownika.docx")
    a = p.paragraphs
    _kropki(a[0], z.get("kierownik_imie_nazwisko", ""), z.get("miejscowosc_sporzadzenia", ""),
            z.get("data_sporzadzenia", ""))
    _kropki(a[7], dane.get("nazwa_inwestycji", ""), _inwestorzy(dane))
    ulica = " ".join(x for x in (adres.get("ulica", ""), adres.get("nr_domu", "")) if x)
    _kropki(a[9], "przy " + ulica if ulica else "", adres.get("miejscowosc", ""),
            adres.get("nr_dzialki", ""))
    _kropki(a[10], "", z.get("pozwolenie_znak", ""), z.get("pozwolenie_data", ""))
    _kropki(a[13], z.get("odstepstwa", ""))
    for numer, pole in ((17, "ep"), (18, "u_sciany"), (19, "u_dach"),
                        (20, "u_podloga"), (21, "u_okna"), (22, "u_drzwi")):
        _kropki(a[numer], z.get(pole, ""))
    if z.get("zgodnosc") == "zgodnie":
        _dopisz(a[11], " [WYBRANO]")
    elif z.get("zgodnosc") == "odstepstwa":
        _dopisz(a[12], " [WYBRANO]")
    else:
        _dopisz(a[11], " [WYBIERZ WARIANT]")
    if z.get("teren_doprowadzony") == "nie":
        _zamien(a[14], "teren budowy został", "teren budowy nie został")
    elif not z.get("teren_doprowadzony"):
        _dopisz(a[14], " [POTWIERDŹ TAK/NIE]")
    if z.get("tereny_przylegle") in {"tak", "nie", "nie_dotyczy"}:
        _zamien(a[15], "tak / nie", z["tereny_przylegle"].replace("_", " "))
    if z.get("pomiary_powierzchni") == "nie":
        _zamien(a[16], "dokonałem pomiarów", "nie dokonałem pomiarów")
    elif not z.get("pomiary_powierzchni"):
        _dopisz(a[16], " [POTWIERDŹ TAK/NIE]")
    if z.get("obiekt_nadaje_sie") == "nie":
        _zamien(a[23], "obiekt nadaje się", "obiekt nie nadaje się")
    elif not z.get("obiekt_nadaje_sie"):
        _dopisz(a[23], " [POTWIERDŹ TAK/NIE]")
    _kropki(p.tables[0].cell(0, 0).paragraphs[0], z.get("kierownik_adres", ""))
    return _zapisz(p)


def zawiadomienie(dane: dict) -> bytes:
    z = dane.get("zakonczenie", {})
    adres = dane.get("adres_inwestycji", {})
    p = Document(SZABLONY / "zawiadomienie_o_zakonczeniu.docx")
    a = p.paragraphs
    # Rok w sygnaturze z pustego wzoru nie jest numerem konkretnej sprawy.
    _zamien(a[0], "2020", "........")
    _kropki(a[2], _inwestorzy(dane), " ".join(x for x in (z.get("miejscowosc_sporzadzenia", ""),
                                                       z.get("data_sporzadzenia", "")) if x))
    _kropki(a[4], "; ".join(_adres(i) for i in dane["inwestorzy"] if _adres(i)))
    _kropki(a[6], z.get("telefon_inwestora", ""))
    _kropki(a[13], dane.get("nazwa_inwestycji", ""))
    ulica = re.sub(r"^ul\.?\s*", "", adres.get("ulica", ""), flags=re.IGNORECASE)
    _kropki(a[15], adres.get("miejscowosc", ""), ulica)
    if ulica:
        _zamien(a[15], "przy ul ", "przy ul. ")
    _kropki(a[16], adres.get("nr_dzialki", ""), z.get("pozwolenie_znak", ""),
            z.get("pozwolenie_data", ""), z.get("pozwolenie_organ", ""))
    _kropki(a[17], z.get("liczba_budynkow", ""), z.get("nr_budynku_sasiada", ""))
    if z.get("odbior_czesciowy"):
        _zamien(a[18], "tak / nie", z["odbior_czesciowy"])
    _kropki(a[19], z.get("pozwolenie_uzytkowanie_data", ""), z.get("pozwolenie_uzytkowanie_znak", ""))
    _kropki(a[20], z.get("data_zgloszenia_rozpoczecia", ""))
    _oznacz(a[21], z.get("rodzaje_obiektu", []), {
        "nowy": "nowy", "rozbudowa": "rozbudowa", "nadbudowa": "nadbudowa",
        "przebudowa": "przebudowa", "zmiana_uzytkowania": "zmiana sposobu użytkowania",
        "adaptacja": "adaptacja", "mieszkalny": "mieszkalny", "niemieszkalny": "niemieszkalny"})
    pola = {25: ("kubatura_mieszkalna", "kubatura_innego"),
            26: ("pow_zabudowy_mieszkalna", "pow_zabudowy_innego"),
            27: ("pow_uzytkowa_mieszkalna", "pow_uzytkowa_innego"),
            28: ("kondygnacje_nadziemne_innego",),
            29: ("pow_mieszkania", "kondygnacje_podziemne_innego"),
            30: ("pow_uslugowa",), 31: ("pow_garazu",), 32: ("pow_calkowita",),
            33: ("liczba_izb",), 34: ("liczba_mieszkan",),
            35: ("kondygnacje_nadziemne",), 36: ("kondygnacje_podziemne",)}
    for numer, nazwy in pola.items():
        _kropki(a[numer], *(z.get(nazwa, "") for nazwa in nazwy))
    _oznacz(a[40], z.get("forma_budownictwa", ""), {
        "indywidualne_wlasne": "indywidualne (inne niż sprzedaż lub wynajem)",
        "indywidualne_sprzedaz": "indywidualne (na sprzedaż lub wynajem)",
        "spoldzielcze": "spółdzielcze", "sprzedaz_wynajem": "na sprzedaż lub wynajem",
        "komunalne": "komunalne", "spoleczne_czynszowe": "społeczne czynszowe",
        "zakladowe": "zakładowe"})
    _oznacz(a[43], z.get("technologia", ""), {
        "tradycyjna": "tradycyjna udoskonalona", "wielkoplytowa": "wielkopłytowa",
        "wielkoblokowa": "wielkoblokowa", "monolityczna": "monolityczna",
        "kanadyjska": "kanadyjska", "inna": "inna"})
    _oznacz(a[44], z.get("rodzaj_scian", ""), {
        "ognioodporne": "ognioodporne", "nieognioodporne": "nieognioodporne", "nro": "NRO"})
    _oznacz(a[46], z.get("kanalizacja", ""), {"miejska": "miejskiej", "wiejska": "wiejskiej"})
    _oznacz(a[47], z.get("kanalizacja", ""), {"zbiornik": "bezodpływowego szczelnego zbiornika",
                                                   "oczyszczalnia": "oczyszczalni przydomowej"})
    _kropki(a[47], z.get("zbiornik_pojemnosc", ""))
    _oznacz(a[50], z.get("pobor_wody", ""), {"studnia_kopana": "kopana", "studnia_wiercona": "wiercona"})
    _oznacz(a[53], z.get("pobor_wody", ""), {
        "wodociag_miejski": "miejski", "wodociag_wiejski": "wiejski",
        "wodociag_zagrodowy": "zagrodowy", "wodociag_indywidualny": "indywidualny"})
    _oznacz(a[55], z.get("ogrzewanie", ""), {
        "olejowe": "olejowe", "gazowe": "gazowe", "opal_staly": "opał stały",
        "pompa_ciepla": "pompa ciepła", "inne": "inne"})
    if z.get("ogrzewanie") == "inne" and z.get("inne_ogrzewanie"):
        _dopisz(a[55], " – " + z["inne_ogrzewanie"])
    if z.get("gaz_sieciowy"):
        _zamien(a[56], "tak / nie", z["gaz_sieciowy"])
    zalaczniki = {59: "dziennik_budowy", 60: "oswiadczenie_kierownika",
                  63: "tereny_przylegle", 65: "przewody_kominowe", 66: "instalacja_elektryczna",
                  67: "instalacja_piorunochronna", 68: "instalacja_gazowa",
                  69: "urzadzenia_ochrony_srodowiska", 70: "instalacje_technologiczne",
                  73: "przylacze_energetyczne", 74: "przylacze_wodociagowe",
                  75: "przylacze_kanalizacyjne", 76: "przylacze_gazowe",
                  77: "przylacze_cieplownicze", 79: "inwentaryzacja_geodezyjna",
                  80: "pozwolenie_na_budowe", 81: "rysunki_ze_zmianami",
                  82: "opis_zmian", 83: "projekt_do_wgladu"}
    for numer, nazwa in zalaczniki.items():
        if nazwa in z.get("zalaczniki", []):
            _dopisz(a[numer], " [x]")
    if z.get("inne_zalaczniki"):
        _dopisz(a[78], " " + z["inne_zalaczniki"])
    return _zapisz(p)
