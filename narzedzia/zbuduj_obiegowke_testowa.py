"""Tworzy obiegówki testowe z fikcyjnymi, kompletnymi danymi (układ jak w obiegówce biura).

Uruchomienie: .venv/bin/python narzedzia/zbuduj_obiegowke_testowa.py
Bazą jest zanonimizowana obiegówka tests/fixtures/obiegowka_przyklad.docx. W każdej tabeli podmieniane są
wartości pól (nazwa, adres inwestycji, inwestor, data) – tak samo na wszystkich stronach.
"""
import copy
import sys
from pathlib import Path

import docx

KATALOG = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KATALOG))
from app.parser_obiegowki import _komorki_wiersza, _rozpoznaj_etykiete  # noqa: E402

BAZA = KATALOG / "tests" / "fixtures" / "obiegowka_przyklad.docx"

WARIANTY = {
    # Jeden inwestor z numerem lokalu, działka w innej miejscowości niż adres inwestora, adres działki z kodem.
    "obiegowka_testowa_pelna.docx": {
        "nazwa_inwestycji": ["„ROZBUDOWA BUDYNKU MIESZKALNEGO JEDNORODZINNEGO O CZĘŚĆ GARAŻOWĄ "
                             "WRAZ Z PRZEBUDOWĄ WEWNĘTRZNĄ”"],
        "adres_inwestycji": ["dz. ewid. 245/7", "121701_1.0005.245/7", "ul. Słoneczna 8, 34-500 Zakopane"],
        "inwestor": ["Krzysztof Testowy", "ul. Przykładowa 12/4", "30-001 Kraków"],
        "kategoria_obiektu": ["I"],
        "data_opracowania": ["MARZEC 2026"],
    },
    # Dwoje inwestorów (wspólny adres), działka pod adresem inwestorów.
    "obiegowka_testowa_dwoje_inwestorow.docx": {
        "nazwa_inwestycji": ["„BUDOWA BUDYNKU MIESZKALNEGO JEDNORODZINNEGO WRAZ Z INSTALACJAMI WEWNĘTRZNYMI”"],
        "adres_inwestycji": ["dz. ewid. 1520/3", "121112_4.0001.1520/3", "Rabka-Zdrój ul. Leśna 3"],
        "inwestor": ["Marek Przykładny i Ewa Przykładna", "ul. Leśna 3", "34-700 Rabka-Zdrój"],
        "kategoria_obiektu": ["I"],
        "data_opracowania": ["KWIECIEŃ 2026"],
    },
}


def ustaw_linie(komorka, linie: list[str]):
    """Zastępuje akapity komórki liniami tekstu, zachowując formatowanie istniejących akapitów."""
    akapity = [p for p in komorka.paragraphs if p.text.strip()] or komorka.paragraphs[:1]
    zostaja = {id(p._p) for p in akapity}
    for p in komorka.paragraphs:
        if id(p._p) not in zostaja:
            p._p.getparent().remove(p._p)
    while len(akapity) < len(linie):
        nowy = copy.deepcopy(akapity[-1]._p)
        akapity[-1]._p.addnext(nowy)
        akapity.append(docx.text.paragraph.Paragraph(nowy, akapity[-1]._parent))
    for p in akapity[len(linie):]:
        p._p.getparent().remove(p._p)
    for p, tekst in zip(akapity, linie):
        runy = [r for r in p.runs if r.text]
        runy[0].text = tekst
        for r in runy[1:]:
            r.text = ""


def zbuduj(wartosci: dict, cel: Path):
    d = docx.Document(BAZA)
    for tabela in d.tables:
        for wiersz in tabela.rows:
            komorki = _komorki_wiersza(wiersz)
            for nr, komorka in enumerate(komorki[:-1]):
                pole = _rozpoznaj_etykiete(komorka.text.replace("\n", " "))
                if pole in wartosci:
                    ustaw_linie(komorki[nr + 1], wartosci[pole])
    d.core_properties.title = "Obiegówka testowa (dane fikcyjne)"
    d.save(cel)


if __name__ == "__main__":
    cel_katalog = Path(sys.argv[1]) if len(sys.argv) > 1 else KATALOG / "tests" / "fixtures"
    for nazwa, wartosci in WARIANTY.items():
        zbuduj(wartosci, cel_katalog / nazwa)
        print("Zapisano", cel_katalog / nazwa)
