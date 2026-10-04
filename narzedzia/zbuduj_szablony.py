"""Buduje szablony biura (szablony/*.docx) od zera, wzorowane na dokumentach FILAR-USŁUGI.

Uruchomienie: .venv/bin/python narzedzia/zbuduj_szablony.py
Znaczniki {{ ... }} wypełnia app/generator.py (docxtpl). Grafika (logo w nagłówku) zostanie dodana później.
"""
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SZABLONY = Path(__file__).resolve().parent.parent / "szablony"
SZARE_TLO = "D9D9D9"
STOPKA_BIURA = ("FILAR-USŁUGI  adres: ul. Sądecka 33h/11, 34-700 Rabka-Zdrój  "
                "tel.: 695 770 613, 795 130 812  e-mail: biurofilarmf@gmail.com")


# ---------------------------------------------------------------- pomocnicze

def nowy_dokument(czcionka: str, rozmiar: float, marginesy: tuple[float, float, float, float]):
    """marginesy w cm: (góra, prawy, dół, lewy)."""
    d = docx.Document()
    styl = d.styles["Normal"]
    styl.font.name = czcionka
    styl.font.size = Pt(rozmiar)
    styl.element.rPr.rFonts.set(qn("w:eastAsia"), czcionka)
    styl.paragraph_format.space_after = Pt(0)
    styl.paragraph_format.line_spacing = 1.0
    sekcja = d.sections[0]
    sekcja.page_width, sekcja.page_height = Cm(21.0), Cm(29.7)
    sekcja.top_margin, sekcja.right_margin, sekcja.bottom_margin, sekcja.left_margin = (Cm(m) for m in marginesy)
    return d


def akapit(miejsce, *fragmenty, wyrownanie=None, przed=0, po=0, tlo=None, wciecie=None, rozmiar=None,
           pogrubienie=False):
    """fragmenty: teksty lub krotki (tekst, {"b": bool, "sup": bool, "size": pt}); "\\n" = złamanie linii."""
    p = miejsce.add_paragraph()
    if wyrownanie is not None:
        p.alignment = wyrownanie
    p.paragraph_format.space_before = Pt(przed)
    p.paragraph_format.space_after = Pt(po)
    if wciecie is not None:
        p.paragraph_format.left_indent = Cm(wciecie)
    if tlo:
        cieniuj(p._p.get_or_add_pPr(), tlo)
    for fragment in fragmenty:
        tekst, opcje = (fragment, {}) if isinstance(fragment, str) else fragment
        for nr, linia in enumerate(tekst.split("\n")):
            if nr:
                p.add_run().add_break(WD_BREAK.LINE)
            if not linia:
                continue
            r = p.add_run(linia)
            r.bold = opcje.get("b", pogrubienie) or None
            r.font.superscript = opcje.get("sup") or None
            if opcje.get("size") or rozmiar:
                r.font.size = Pt(opcje.get("size") or rozmiar)
    return p


def cieniuj(element_pr, kolor: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), kolor)
    element_pr.append(shd)


def stopka(d, tekst: str, rozmiar: float = 7.5):
    p = d.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(tekst)
    r.font.size = Pt(rozmiar)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


# ---------------------------------------------------------------- PB-5

def pb5():
    d = nowy_dokument("Times New Roman", 10.5, (0.75, 1.8, 1.25, 1.8))

    # Nagłówek formularza na szarym tle
    tabela = d.add_table(rows=1, cols=1)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    komorka = tabela.cell(0, 0)
    cieniuj(komorka._tc.get_or_add_tcPr(), SZARE_TLO)
    komorka.paragraphs[0]._p.getparent().remove(komorka.paragraphs[0]._p)
    akapit(komorka, "OŚWIADCZENIE", wyrownanie=WD_ALIGN_PARAGRAPH.CENTER, przed=4, rozmiar=13, pogrubienie=True)
    akapit(komorka, "o posiadanym prawie do dysponowania nieruchomością na cele budowlane",
           wyrownanie=WD_ALIGN_PARAGRAPH.CENTER, rozmiar=11, pogrubienie=True)
    akapit(komorka, "(PB-5)", wyrownanie=WD_ALIGN_PARAGRAPH.CENTER, po=6, rozmiar=11, pogrubienie=True)
    akapit(komorka, ("Podstawa prawna", {"b": True}),
           ": Art. 32 ust. 4 pkt 2 ustawy z dnia 7 lipca 1994 r. – Prawo budowlane "
           "(Dz. U. z 2020 r. poz. 1333, z późn. zm.).", rozmiar=7.5, po=3, wciecie=0.3)
    akapit(komorka, ("Dodatkowe informacje: ", {"b": True}),
           "Prawo do dysponowania nieruchomością na cele budowlane jest to tytuł prawny wynikający z prawa "
           "własności, użytkowania wieczystego, zarządu, ograniczonego prawa rzeczowego albo stosunku "
           "zobowiązaniowego, przewidującego uprawnienia do wykonywania robót budowlanych.",
           wyrownanie=WD_ALIGN_PARAGRAPH.JUSTIFY, rozmiar=7.5, po=3, wciecie=0.3)
    akapit(komorka, "W przypadku, gdy do złożenia oświadczenia zobowiązanych jest kilka osób, każda z tych osób "
                    "składa oświadczenie oddzielnie na osobnym formularzu.",
           wyrownanie=WD_ALIGN_PARAGRAPH.JUSTIFY, rozmiar=7.5, po=4, wciecie=0.3)

    def naglowek_sekcji(*fragmenty):
        akapit(d, *fragmenty, przed=10, po=6, tlo=SZARE_TLO, pogrubienie=True)

    def linia(*fragmenty):
        akapit(d, *fragmenty, po=6, wciecie=0.5)

    naglowek_sekcji("1. DANE INWESTORA")
    linia("Imię i nazwisko lub nazwa: {{ inwestor.imie_nazwisko }}")
    linia("Kraj: {{ inwestor.kraj }} Województwo: {{ inwestor.wojewodztwo }} "
          "Powiat: {{ inwestor.powiat }} Gmina: {{ inwestor.gmina }}")
    linia("Ulica: {{ inwestor.ulica }} Nr domu: {{ inwestor.nr_domu }} Nr lokalu: {{ inwestor.nr_lokalu }}")
    linia("Miejscowość: {{ inwestor.miejscowosc }} Kod pocztowy: {{ inwestor.kod_pocztowy }} "
          "Poczta: {{ inwestor.poczta }}")

    naglowek_sekcji("2. DANE OSOBY UPOWAŻNIONEJ DO ZŁOŻENIA OŚWIADCZENIA W IMIENIU INWESTORA", ("1)", {"sup": True}))
    linia("Imię i nazwisko lub nazwa: {{ upowazniony.imie_nazwisko }}")
    linia("Kraj: {{ upowazniony.kraj }} Województwo: {{ upowazniony.wojewodztwo }}")
    linia("Powiat: {{ upowazniony.powiat }} Gmina: {{ upowazniony.gmina }}")
    linia("Ulica: {{ upowazniony.ulica }} Nr domu: {{ upowazniony.nr_domu }} Nr lokalu: {{ upowazniony.nr_lokalu }}")
    linia("Miejscowość: {{ upowazniony.miejscowosc }} Kod pocztowy: {{ upowazniony.kod_pocztowy }} "
          "Poczta: {{ upowazniony.poczta }}")

    naglowek_sekcji("3. DANE NIERUCHOMOŚCI", ("2)", {"sup": True}))
    linia("Województwo: {{ nieruchomosc.wojewodztwo }}")
    linia("Powiat: {{ nieruchomosc.powiat }} Gmina: {{ nieruchomosc.gmina }}")
    linia("Ulica: {{ nieruchomosc.ulica }} Nr domu: {{ nieruchomosc.nr_domu }}")
    linia("Miejscowość: {{ nieruchomosc.miejscowosc }} Kod pocztowy: {{ nieruchomosc.kod_pocztowy }}")
    linia("Identyfikator działki ewidencyjnej", ("3)", {"sup": True}), ": {{ nieruchomosc.identyfikator_dzialki }}")
    akapit(d, "Liczba stron zawierających dane o kolejnych nieruchomościach (załączanych do oświadczenia): ……",
           po=6, wciecie=0.5, rozmiar=9.5)

    akapit(d, "Po zapoznaniu się z art. 32 ust. 4 pkt 2 oraz art. 3 pkt 11 ustawy z dnia 7 lipca 1994 r. – Prawo "
              "budowlane oświadczam, że posiadam prawo do dysponowania nieruchomością (nieruchomościami) na cele "
              "budowlane określoną (określonymi) w pkt 3 tego oświadczenia.",
           wyrownanie=WD_ALIGN_PARAGRAPH.JUSTIFY, przed=10, po=4, tlo=SZARE_TLO, rozmiar=9.5)
    akapit(d, "Jestem świadomy (świadoma) odpowiedzialności karnej za podanie nieprawdy w niniejszym oświadczeniu, "
              "zgodnie z art. 233 ustawy z dnia 6 czerwca 1997 r. – Kodeks karny (Dz. U. z 2020 r. poz. 1444, "
              "z późn. zm.).", wyrownanie=WD_ALIGN_PARAGRAPH.JUSTIFY, po=4, tlo=SZARE_TLO, rozmiar=9.5)

    naglowek_sekcji("4. PODPIS INWESTORA LUB OSOBY UPOWAŻNIONEJ DO ZŁOŻENIA OŚWIADCZENIA W IMIENIU INWESTORA "
                    "I DATA PODPISU")
    akapit(d, "Podpis powinien być czytelny.", po=28, rozmiar=7.5)
    akapit(d, "……………………………………………………………………………………………………")

    akapit(d, "Objaśnienia:", przed=24, po=2, rozmiar=7.5, pogrubienie=True)
    for przypis in [
        "1) Wypełnia się, jeżeli oświadczenie jest składane w imieniu osoby prawnej lub jednostki organizacyjnej "
        "nieposiadającej osobowości prawnej albo oświadczenie w imieniu inwestora składa jego pełnomocnik.",
        "2) W przypadku większej liczby nieruchomości dane kolejnych nieruchomości dodaje się w formularzu albo "
        "zamieszcza na osobnych stronach i dołącza do formularza.",
        "3) W przypadku oświadczenia sporządzanego w postaci papierowej zamiast identyfikatora działki ewidencyjnej "
        "można wskazać obręb ewidencyjny i nr działki ewidencyjnej oraz arkusz mapy, jeżeli występuje.",
    ]:
        akapit(d, przypis, wyrownanie=WD_ALIGN_PARAGRAPH.JUSTIFY, po=1, rozmiar=7.5)

    d.save(SZABLONY / "pb5_oswiadczenie.docx")


# ---------------------------------------------------------------- Pełnomocnictwo

def pelnomocnictwo():
    d = nowy_dokument("Arial Narrow", 12, (2.5, 1.75, 1.5, 2.5))
    sekcja = d.sections[0]
    sekcja.footer_distance = Cm(0.8)

    akapit(d, "Rabka-Zdrój {{ data }}", wyrownanie=WD_ALIGN_PARAGRAPH.RIGHT, po=36)
    akapit(d, "Pełnomocnictwo", wyrownanie=WD_ALIGN_PARAGRAPH.CENTER, po=30, rozmiar=20, pogrubienie=True)

    akapit(d, "{{ ja }} {{ mocodawcy }}", po=6)
    akapit(d, "{{ zamieszkaly }}: {{ adres_mocodawcy }}", po=12)
    akapit(d, "{{ udzielam }} pełnomocnictwa ", ("Bartłomiejowi Filar", {"b": True}), po=6)
    akapit(d, "zam. 34-700 Rabka-Zdrój ul. Sądecka 33h/11 do {{ mnie }} jako {{ inwestora }} w sprawie wszelkich "
              "formalności oraz do występowania w {{ moim }} imieniu przed organami administracji publicznej "
              "związanych z:", wyrownanie=WD_ALIGN_PARAGRAPH.JUSTIFY, po=12)
    akapit(d, "Uzyskaniem wszelkich uzgodnień i decyzji w związku z realizacją inwestycji „{{ nazwa_inwestycji }}”",
           po=12, rozmiar=14, pogrubienie=True)
    akapit(d, "w miejscowości:", po=6)
    akapit(d, "{{ lokalizacja }}", po=18, rozmiar=14, pogrubienie=True)
    akapit(d, "{{ upowazniam }} pełnomocnika do składania wszelkich oświadczeń, wniosków, podań, odwołań, skarg oraz "
              "zażaleń, jak również do odbioru i kwitowania odbioru dokumentów, korespondencji, prawa do zrzeczenia "
              "się od odwołania od wydanej decyzji, oraz do wszystkich innych czynności, jakie okażą się niezbędne "
              "do wykonania niniejszego pełnomocnictwa.", wyrownanie=WD_ALIGN_PARAGRAPH.JUSTIFY, po=48, rozmiar=13)
    akapit(d, "{{ podpisy }}", wyrownanie=WD_ALIGN_PARAGRAPH.RIGHT)

    stopka(d, STOPKA_BIURA)
    d.save(SZABLONY / "pelnomocnictwo.docx")


if __name__ == "__main__":
    SZABLONY.mkdir(exist_ok=True)
    pb5()
    pelnomocnictwo()
    print(f"Zapisano szablony w {SZABLONY}")
